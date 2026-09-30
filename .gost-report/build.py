"""Сборка отчета по лабораторной работе 1 (docs/report.docx).

Верстка повторяет структуру образца: титульный лист, раздел "Выполнение"
с полужирными подзаголовками, код и ответы API моноширинным шрифтом,
скриншоты отчетов pipeline и ссылка на последний запуск.

Запуск: GOST_REPORT_CONFIG=/dev/null gr .gost-report/build.py
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/report.docx"

REPO = "https://github.com/virdxk/infosec-lab1"
RUN = "https://github.com/virdxk/infosec-lab1/actions/runs/36760742483"
SUCCESS_RUNS = "https://github.com/virdxk/infosec-lab1/actions/workflows/ci.yml?query=is%3Asuccess"

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
MONO_SIZE = Pt(9)
CODE_MAX_COLS = 88
LINK_COLOR = RGBColor(0x11, 0x55, 0xCC)

doc = Document()
figure_no = 0
table_no = 0


# ---------- низкоуровневые помощники ----------

def set_font(run, name=BODY_FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def paragraph(align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, indent=True,
              keep_next=False, line=1.15):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.alignment = align
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.first_line_indent = Cm(1.25) if indent else Cm(0)
    fmt.keep_with_next = keep_next
    fmt.widow_control = True
    return p


def add_link(p, url, label=None, size=None):
    rel = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    link.append(run)
    p._p.append(link)
    from docx.text.run import Run
    r = Run(run, p)
    r.text = label or url
    set_font(r, size=size, color=LINK_COLOR)
    r.underline = True
    return r


INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))")


def add_inline(p, text, size=None):
    """Мини-разметка: `код`, **жирный**, [текст](url)."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("`"):
            set_font(p.add_run(part[1:-1]), MONO_FONT, size=Pt(10.5) if size is None else size)
        elif part.startswith("**"):
            set_font(p.add_run(part[2:-2]), size=size, bold=True)
        elif part.startswith("["):
            label, url = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            add_link(p, url, label, size=size)
        else:
            set_font(p.add_run(part), size=size)


def shade(element_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element_pr.append(shd)


def cell_borders(cell, color="BFBFBF", size=4):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:color"), color)
        borders.append(b)
    tcpr.append(borders)


def cell_margins(table, top=80, bottom=80, left=140, right=140):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def page_field(run):
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._r.append(instr)


# ---------- элементы отчета ----------

def h1(text):
    p = paragraph(align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=10, indent=False, keep_next=True)
    set_font(p.add_run(text), size=Pt(16), bold=True)


def h2(text):
    p = paragraph(align=WD_ALIGN_PARAGRAPH.LEFT, before=14, after=6, indent=False, keep_next=True)
    set_font(p.add_run(text), size=Pt(14), bold=True)


def h3(text):
    p = paragraph(align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=4, indent=False, keep_next=True)
    set_font(p.add_run(text), size=Pt(12), bold=True)


def text(value, keep_next=False):
    p = paragraph(keep_next=keep_next)
    add_inline(p, value)
    return p


def bullets(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        fmt = p.paragraph_format
        fmt.space_after = Pt(3)
        fmt.line_spacing = 1.15
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(p, item)


def link_line(url, label=None):
    p = paragraph(align=WD_ALIGN_PARAGRAPH.LEFT, indent=False)
    add_link(p, url, label)


def code(source, label=None):
    """Блок кода: одна ячейка с серым фоном, Consolas 9 pt, без переносов."""
    lines = source.strip("\n").splitlines()
    too_long = [ln for ln in lines if len(ln) > CODE_MAX_COLS]
    if too_long:
        raise ValueError(f"строка кода длиннее {CODE_MAX_COLS} символов: {too_long[0]!r}")
    if label:
        p = paragraph(align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=2, indent=False, keep_next=True)
        set_font(p.add_run(label), size=Pt(10), italic=True, color=RGBColor(0x55, 0x55, 0x55))
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(table)
    cell = table.cell(0, 0)
    cell_borders(cell)
    shade(cell._tc.get_or_add_tcPr(), "F4F4F4")
    if len(lines) <= 40:
        no_split(table.rows[0])
    first = True
    for ln in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        fmt = p.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0
        fmt.first_line_indent = Cm(0)
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(ln if ln else " "), MONO_FONT, size=MONO_SIZE)
    paragraph(after=0, indent=False, line=0.6)


def table(rows, widths, caption, mono_cols=(), size=Pt(10.5)):
    global table_no
    table_no += 1
    p = paragraph(align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3, indent=False, keep_next=True)
    set_font(p.add_run(f"Таблица {table_no} - {caption}"), size=Pt(11))
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t, top=40, bottom=40, left=90, right=90)
    for r, row in enumerate(rows):
        no_split(t.rows[r])
        if r == 0:
            repeat_header(t.rows[r])
        for c, value in enumerate(row):
            cell = t.cell(r, c)
            cell.width = Cm(widths[c])
            cell_borders(cell, color="808080")
            if r == 0:
                shade(cell._tc.get_or_add_tcPr(), "E8E8E8")
            cp = cell.paragraphs[0]
            fmt = cp.paragraph_format
            fmt.space_after = Pt(0)
            fmt.line_spacing = 1.0
            fmt.first_line_indent = Cm(0)
            fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.keep_with_next = r < len(rows) - 1
            if r == 0:
                set_font(cp.add_run(value), size=size, bold=True)
            elif c in mono_cols:
                set_font(cp.add_run(value), MONO_FONT, size=Pt(9))
            else:
                add_inline(cp, value, size=size)
    paragraph(after=0, indent=False, line=0.8)


def figure(path, caption, width_cm=15.0, border=False):
    global figure_no
    figure_no += 1
    p = paragraph(align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=4, indent=False, keep_next=True, line=1.0)
    if border:
        pbdr = OxmlElement("w:pBdr")
        for side in ("top", "left", "bottom", "right"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "1")
            b.set(qn("w:color"), "A0A0A0")
            pbdr.append(b)
        p._p.get_or_add_pPr().append(pbdr)
    p.add_run().add_picture(str(ROOT / path), width=Cm(width_cm))
    p = paragraph(align=WD_ALIGN_PARAGRAPH.CENTER, after=10, indent=False)
    set_font(p.add_run(f"Рисунок {figure_no} - {caption}"), size=Pt(11))


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------- документ ----------

normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
bullet_style = doc.styles["List Bullet"]
bullet_style.font.name = BODY_FONT
bullet_style.font.size = Pt(12)

section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.left_margin, section.right_margin = Cm(2.5), Cm(1.5)
section.top_margin, section.bottom_margin = Cm(2), Cm(2)
section.different_first_page_header_footer = True

footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer_p.add_run()
set_font(run, size=Pt(11))
page_field(run)

first_footer = section.first_page_footer.paragraphs[0]
first_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(first_footer.add_run("Санкт-Петербург"), size=Pt(14))
set_font(section.first_page_footer.add_paragraph().add_run("2026"), size=Pt(14))
section.first_page_footer.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Титульный лист
for line in ("Федеральное государственное автономное",
             "образовательное учреждение высшего образования",
             "«Национальный исследовательский университет ИТМО»"):
    p = paragraph(align=WD_ALIGN_PARAGRAPH.CENTER, after=0, indent=False)
    set_font(p.add_run(line), size=Pt(14))

p = paragraph(align=WD_ALIGN_PARAGRAPH.CENTER, before=190, after=14, indent=False)
set_font(p.add_run("Информационная безопасность"), size=Pt(16), bold=True)
p = paragraph(align=WD_ALIGN_PARAGRAPH.CENTER, after=14, indent=False)
set_font(p.add_run("Лабораторная работа № 1"), size=Pt(16), bold=True)
p = paragraph(align=WD_ALIGN_PARAGRAPH.CENTER, after=0, indent=False)
set_font(p.add_run("Разработка защищенного REST API\nс интеграцией в CI/CD"), size=Pt(16), bold=True)

p = paragraph(align=WD_ALIGN_PARAGRAPH.RIGHT, before=150, after=4, indent=False)
set_font(p.add_run("Выполнил: Елисеев Константин Иванович"), size=Pt(14))
p = paragraph(align=WD_ALIGN_PARAGRAPH.RIGHT, after=0, indent=False)
set_font(p.add_run("Группа: P3308"), size=Pt(14))
page_break()

# Выполнение
h1("Выполнение")

h2("Ссылка на GitHub репозиторий")
link_line(REPO)

h2("Описание проекта")
text("Небольшой REST API на Python: вход по логину и паролю с выдачей JWT, чтение и "
     "создание постов. Стек: Flask, SQLite через SQLAlchemy, PyJWT, bcrypt.")
table([
    ["Файл", "Что делает"],
    ["`app/__init__.py`", "Создание приложения, проверка длины секрета, подключение blueprints"],
    ["`app/config.py`", "Настройки: секрет и адрес БД из `.env`, срок токена, стоимость bcrypt"],
    ["`app/auth.py`", "Вход по логину и паролю, выдача токена"],
    ["`app/security.py`", "Выпуск и проверка JWT"],
    ["`app/middleware.py`", "Проверка JWT перед маршрутами `/api/*`"],
    ["`app/api.py`", "Чтение и создание постов"],
    ["`app/models.py`", "Модели User и Post, хэширование паролей bcrypt"],
    ["`app/sanitize.py`", "Экранирование пользовательских данных в ответах"],
    ["`seed.py`", "Демо-пользователи alice и bob, два поста"],
], widths=[4.5, 12.3], caption="Структура проекта")

h3("Запуск")
code("""
python3 -m venv .venv
source .venv/bin/activate
pip install -r requirements-dev.txt
cp .env.example .env   # вписать JWT_SECRET, сгенерировать командой из файла
python seed.py         # пользователи alice и bob, два поста
python wsgi.py         # http://127.0.0.1:5000
""")
text("В `.env` задаются только `JWT_SECRET` и `DATABASE_URL`. Срок жизни токена (3600 с) "
     "и стоимость bcrypt (12) зафиксированы в `app/config.py`.")

h2("Описание API")
table([
    ["Метод и адрес", "Назначение", "Токен", "Успешный ответ"],
    ["`POST /auth/login`", "Вход, тело `{\"username\", \"password\"}`", "не нужен",
     "`200 {\"access_token\"}`"],
    ["`GET /api/data`", "Список постов", "нужен", "`200 [ {id, title, body, author} ]`"],
    ["`POST /api/posts`", "Создать пост, тело `{\"title\", \"body\"}`", "нужен",
     "`201 {id, title, body, author}`"],
], widths=[3.6, 5.6, 1.8, 5.8], caption="Эндпоинты API")
text("Токен передается в заголовке `Authorization: Bearer <token>`. Ошибки: 400 при неполном "
     "теле запроса, 401 при неверном пароле или отсутствующем, поддельном либо просроченном "
     "токене. В примерах `$BASE_URL` - адрес сервера, `$TOKEN` - токен из ответа на вход.")

h3("POST /auth/login")
code(r"""
curl -s -X POST "$BASE_URL/auth/login" \
  -H 'Content-Type: application/json' \
  -d '{"username": "alice", "password": "Al1ce-Str0ng-Pass!"}'

HTTP 200
{
  "access_token": "<JWT>"
}
""")

h3("GET /api/data")
code(r"""
curl -s "$BASE_URL/api/data" -H "Authorization: Bearer $TOKEN"

HTTP 200
[
  {
    "author": "alice",
    "body": "First post.",
    "id": 1,
    "title": "Welcome"
  },
  {
    "author": "bob",
    "body": "Tokens expire in one hour.",
    "id": 2,
    "title": "Notes on JWT"
  }
]
""")

h3("POST /api/posts")
code(r"""
curl -s -X POST "$BASE_URL/api/posts" \
  -H "Authorization: Bearer $TOKEN" \
  -H 'Content-Type: application/json' \
  -d '{"title": "Hello", "body": "World"}'

HTTP 201
{
  "author": "alice",
  "body": "World",
  "id": 3,
  "title": "Hello"
}
""")

h2("Описание реализованных мер защиты")

h3("Защита от SQL-инъекций")
text("Запросы к базе строятся через ORM SQLAlchemy. Логин передается в SQL связанным "
     "параметром, поэтому ввод `' OR 1=1 --` остается обычной строкой логина. SQL из строк "
     "нигде не склеивается.", keep_next=True)
code("""
with current_app.session_factory() as session:
    user = session.scalar(select(User).where(User.username == username))

if user is None or not user.check_password(password):
    return jsonify({"error": "invalid credentials"}), 401
""", label="app/auth.py")

h3("Защита от XSS")
text("Перед отдачей клиенту заголовок, текст поста и имя автора экранируются функцией "
     "`markupsafe.escape`. В базе хранится исходный текст, экранирование применяется к "
     "каждому ответу.", keep_next=True)
code("""
def serialize_post(post: Post) -> dict:
    return {
        "id": post.id,
        "title": str(escape(post.title)),
        "body": str(escape(post.body)),
        "author": str(escape(post.author.username)),
    }
""", label="app/sanitize.py")
code(r"""
curl -s -X POST "$BASE_URL/api/posts" \
  -H "Authorization: Bearer $TOKEN" \
  -H 'Content-Type: application/json' \
  -d '{"title": "XSS", "body": "<script>alert(1)</script>"}'

HTTP 201
{
  "author": "alice",
  "body": "&lt;script&gt;alert(1)&lt;/script&gt;",
  "id": 4,
  "title": "XSS"
}
""", label="Проверка")

h3("Хранение паролей: bcrypt")
text("В базе хранятся только bcrypt-хэши. Соль генерируется случайно для каждого пароля и "
     "хранится внутри хэша, стоимость - 12 раундов. Пароли длиннее 72 байт bcrypt не "
     "принимает, поэтому для них проверка сразу возвращает отказ.", keep_next=True)
code("""
def set_password(self, password: str, rounds: int) -> None:
    salt = bcrypt.gensalt(rounds=rounds)
    self.password_hash = bcrypt.hashpw(password.encode(), salt).decode()

def check_password(self, password: str) -> bool:
    password_bytes = password.encode()
    # bcrypt raises ValueError on passwords longer than 72 bytes
    if len(password_bytes) > 72:
        return False
    return bcrypt.checkpw(password_bytes, self.password_hash.encode())
""", label="app/models.py")
text("На неверный пароль и на несуществующего пользователя ответ одинаковый: "
     "`401 {\"error\": \"invalid credentials\"}`.")

h3("Аутентификация: JWT")
text("После входа выдается токен HS256 с полями `sub` (ID пользователя), `iat` и `exp` "
     "(срок 1 час). Секрет берется из `.env`; если он короче 32 символов, приложение не "
     "запускается. При проверке алгоритм задан сервером, поэтому токен с `alg: none` "
     "отклоняется.", keep_next=True)
code("""
def issue_token(user_id: int) -> str:
    now = datetime.now(timezone.utc)
    payload = {
        "sub": str(user_id),
        "iat": now,
        "exp": now + timedelta(seconds=current_app.config["JWT_TTL_SECONDS"]),
    }
    return jwt.encode(payload, current_app.config["JWT_SECRET"], algorithm=ALGORITHM)


def decode_token(token: str) -> dict:
    return jwt.decode(
        token,
        current_app.config["JWT_SECRET"],
        algorithms=[ALGORITHM],
        options={"require": ["sub", "iat", "exp"]},
    )
""", label="app/security.py")
text("Middleware подключено ко всем маршрутам `/api/*` через хук `before_request`. Оно "
     "проверяет заголовок, подпись и срок токена, находит пользователя в базе и только "
     "тогда пропускает запрос, иначе отвечает 401. Автор нового поста берется из токена, "
     "а не из тела запроса.", keep_next=True)
code("""
def authenticate_request():
    header = request.headers.get("Authorization", "")
    if not header.startswith("Bearer "):
        return jsonify({"error": "missing token"}), 401

    try:
        claims = decode_token(header.removeprefix("Bearer "))
    except jwt.InvalidTokenError:
        return jsonify({"error": "invalid token"}), 401

    with current_app.session_factory() as session:
        user = session.get(User, int(claims["sub"]))
    if user is None:
        return jsonify({"error": "invalid token"}), 401

    g.user_id = user.id
    return None
""", label="app/middleware.py")

h2("Проверка API через curl")
text("Сервер запущен локально на чистой базе после `python seed.py`. Полный протокол с "
     "командами и ответами - файл `docs/curl-session.md` в репозитории.", keep_next=True)
table([
    ["№", "Сценарий", "Код", "Ответ"],
    ["1", "Вход alice с верным паролем", "200", "access_token"],
    ["2", "Неверный пароль", "401", "invalid credentials"],
    ["3", "SQL-инъекция в логине", "401", "invalid credentials"],
    ["4", "GET /api/data без токена", "401", "missing token"],
    ["5", "Токен без подписи (alg: none)", "401", "invalid token"],
    ["6", "GET /api/data с токеном", "200", "список из двух постов"],
    ["7", "Создание поста", "201", "пост с id 3"],
    ["8", "Пост со <script>", "201", "body экранирован"],
    ["9", "Пост без поля body", "400", "title and body are required"],
], widths=[1.0, 7.0, 1.4, 7.4], caption="Результаты проверки через curl")
code(r"""
curl -s -X POST "$BASE_URL/auth/login" \
  -H 'Content-Type: application/json' \
  -d "{\"username\": \"' OR 1=1 --\", \"password\": \"x\"}"

HTTP 401
{
  "error": "invalid credentials"
}
""", label="Сценарий 3: SQL-инъекция")

h2("CI/CD pipeline")
text("Файл `.github/workflows/ci.yml` запускается на каждый push и pull request. В нем три "
     "независимые задачи, любая найденная проблема завершает pipeline ошибкой.")
table([
    ["Задача", "Что делает", "Когда падает"],
    ["Tests", "`pytest -v`: вход, SQL-инъекция, хэш пароля, доступ без токена, поддельный, "
              "`alg: none` и просроченный токены, создание поста, XSS", "Упал хотя бы один тест"],
    ["SAST (bandit)", "Bandit по `app`, `seed.py`, `wsgi.py`, текстовый отчет в артефакт "
                      "`bandit-report`", "Замечание уровня medium или high"],
    ["SCA (Dependency-Check)", "OWASP Dependency-Check 13.0.0 по всем зависимостям, включая "
     "транзитивные, отчеты HTML и JSON в артефакт `sca-reports`", "Уязвимость с CVSS 7 и выше"],
], widths=[3.4, 8.6, 4.8], caption="Задачи pipeline")
text("Результаты запуска CI #9 (коммит `927bd82`):", keep_next=True)
bullets([
    "Tests: **14 тестов пройдено**;",
    "Bandit: **0 замечаний** (175 строк кода);",
    "OWASP Dependency-Check: **13 зависимостей, 0 уязвимостей**, одно ложное срабатывание "
    "исключено.",
])
text("Исключенное срабатывание - CVE-2025-45770. Оно относится к PHP-библиотеке "
     "`lcobucci/jwt`, а сканер сопоставил его с Python-пакетом PyJWT. Исключение описано в "
     "`.github/dependency-check-suppressions.xml` и ограничено этим пакетом и CVE.")

h2("Отчеты из pipeline")
figure("docs/screenshots/01-actions-run-summary.png",
       "Запуск CI #9: все три задачи завершились успешно, 2 артефакта")
figure("docs/screenshots/02-ci-sast-bandit.png",
       "Отчет Bandit из артефакта bandit-report")
figure("docs/screenshots/03-ci-sca-dependency-check.png",
       "HTML-отчет OWASP Dependency-Check из артефакта sca-reports", width_cm=13.0, border=True)

h2("Последний запуск пайплайна")
text("Проверенный запуск, к которому относятся скриншоты и отчеты:", keep_next=True)
link_line(RUN)
text("Все успешные запуски workflow:", keep_next=True)
link_line(SUCCESS_RUNS)

props = doc.core_properties
props.author = "Елисеев Константин Иванович"
props.title = "Лабораторная работа № 1. Разработка защищенного REST API с интеграцией в CI/CD"
props.last_modified_by = props.author

doc.save(OUT)
print(OUT)
